"""
merge_lichsu.py - Gộp nhiều file Excel Lịch Sử 11 thành một file thống nhất
=============================================================================
Chiến lược gộp:
  - Đọc tất cả file đầu vào (file lớn đã có + file chưa gộp)
  - Căn chỉnh cột (các file có số cột trích dẫn khác nhau)
  - Nối dữ liệu theo thứ tự: không ghi đè nội dung đã có
  - Đánh lại STT liên tục
  - Lưu ra file đích (mặc định LichSu11.xlsx)

Cách dùng:
  # Lần đầu: gộp 3 file nguồn thành 1
  python merge_lichsu.py --files file1.xlsx file2.xlsx file3.xlsx

  # Thêm file mới vào file lớn đã tồn tại (KHÔNG ghi đè)
  python merge_lichsu.py --base LichSu11.xlsx --files file_moi.xlsx

  # Chỉ định thư mục chứa file cần gộp
  python merge_lichsu.py --folder ./pending/ --base LichSu11.xlsx

  # Tuỳ chỉnh output
  python merge_lichsu.py --files a.xlsx b.xlsx --output KetQua.xlsx
"""

import os
import sys
import argparse
import glob
from pathlib import Path
from datetime import datetime

import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ══════════════════════════════════════════════════════════════════════════════
# CẤU HÌNH
# ══════════════════════════════════════════════════════════════════════════════

FIXED_COLS = ["STT", "Chủ đề / Chương", "Bài", "Mục", "Tiểu mục", "Số lượng đoạn trích"]

# Key duy nhất để nhận diện một dòng đã tồn tại (tránh trùng lặp khi gộp)
UNIQUE_KEY_COLS = ["Bài", "Mục", "Tiểu mục"]


# ══════════════════════════════════════════════════════════════════════════════
# ĐỌC VÀ XỬ LÝ DỮ LIỆU
# ══════════════════════════════════════════════════════════════════════════════

def read_excel_safe(path: str) -> pd.DataFrame:
    """Đọc file Excel, giữ nguyên kiểu dữ liệu chuỗi."""
    try:
        df = pd.read_excel(path, dtype=str, header=0)
        df = df.where(pd.notna(df), None)  # NaN → None
        print(f"  ✅ Đọc OK: {path}  →  {len(df)} dòng, {len(df.columns)} cột")
        return df
    except Exception as e:
        print(f"  ❌ Lỗi đọc file {path}: {e}")
        sys.exit(1)


def get_max_trich_num(df: pd.DataFrame) -> int:
    """Lấy số trích dẫn lớn nhất trong DataFrame (từ tên cột)."""
    max_n = 0
    for col in df.columns:
        col_s = str(col)
        if col_s.startswith("Nội dung trích dẫn "):
            try:
                n = int(col_s.replace("Nội dung trích dẫn ", "").strip())
                max_n = max(max_n, n)
            except ValueError:
                pass
    return max_n


def build_unified_columns(dfs: list[pd.DataFrame]) -> list[str]:
    """
    Xây dựng danh sách cột thống nhất:
    - 6 cột cố định đầu
    - Các cột trích dẫn theo số lớn nhất từ tất cả file
    """
    max_n = max(get_max_trich_num(df) for df in dfs)
    trich_cols = []
    for i in range(1, max_n + 1):
        trich_cols.append(f"Nội dung trích dẫn {i}")
        trich_cols.append(f"Trích nguồn {i}")
    return FIXED_COLS + trich_cols


def is_empty(val) -> bool:
    """Kiểm tra một giá trị có thực sự rỗng không (None, nan, chuỗi rỗng)."""
    if val is None:
        return True
    try:
        import math
        if math.isnan(float(val)):
            return True
    except (TypeError, ValueError):
        pass
    return str(val).strip().lower() in ("", "none", "nan")


def align_df_to_columns(df: pd.DataFrame, unified_cols: list[str]) -> pd.DataFrame:
    """
    Căn chỉnh DataFrame về đúng cấu trúc cột thống nhất:
    - Cột thiếu → thêm vào với giá trị None
    - Dùng pd.concat để tránh fragmentation với DataFrame nhiều cột
    """
    missing_cols = {col: pd.Series([None] * len(df), dtype=object)
                    for col in unified_cols if col not in df.columns}
    if missing_cols:
        extra = pd.DataFrame(missing_cols, index=df.index)
        df = pd.concat([df, extra], axis=1)
    return df[unified_cols].copy()


def make_unique_key(row: pd.Series) -> str:
    """Tạo khoá duy nhất từ các cột định danh."""
    parts = []
    for col in UNIQUE_KEY_COLS:
        val = row.get(col, "")
        parts.append("" if is_empty(val) else str(val).strip())
    return "|||".join(parts)


def merge_row_data(existing_row: pd.Series, new_row: pd.Series,
                   unified_cols: list[str]) -> pd.Series:
    """
    Gộp 2 dòng cùng key: ưu tiên giữ dữ liệu đã có trong existing_row.
    Nếu ô nào trong existing_row trống, lấy từ new_row.
    KHÔNG ghi đè nội dung đã có.
    """
    result = existing_row.copy()
    for col in unified_cols:
        if is_empty(result.get(col)):
            new_val = new_row.get(col)
            if not is_empty(new_val):
                result[col] = new_val
    return result


# ══════════════════════════════════════════════════════════════════════════════
# LOGIC GỘP CHÍNH
# ══════════════════════════════════════════════════════════════════════════════

def merge_dataframes(base_df: pd.DataFrame | None,
                     new_dfs: list[pd.DataFrame],
                     unified_cols: list[str]) -> pd.DataFrame:
    """
    Gộp base_df (nếu có) với danh sách new_dfs.
    - Dòng đã tồn tại (cùng key): merge không ghi đè
    - Dòng mới: thêm vào cuối
    """
    # Bắt đầu từ base hoặc rỗng
    if base_df is not None:
        merged_data: dict[str, pd.Series] = {}
        for _, row in base_df.iterrows():
            key = make_unique_key(row)
            if key:
                merged_data[key] = row
        # Giữ thứ tự
        ordered_keys: list[str] = [make_unique_key(row) for _, row in base_df.iterrows()]
        ordered_keys = list(dict.fromkeys(ordered_keys))  # dedup giữ thứ tự
    else:
        merged_data = {}
        ordered_keys = []

    # Gộp từng file mới
    for df in new_dfs:
        added = 0
        merged_count = 0
        for _, row in df.iterrows():
            key = make_unique_key(row)
            if not key:
                continue
            if key in merged_data:
                # Merge không ghi đè
                merged_data[key] = merge_row_data(merged_data[key], row, unified_cols)
                merged_count += 1
            else:
                merged_data[key] = row
                ordered_keys.append(key)
                added += 1
        print(f"    → Thêm mới: {added} dòng | Merge vào dòng đã có: {merged_count} dòng")

    # Tái tạo DataFrame theo đúng thứ tự
    rows = [merged_data[k] for k in ordered_keys if k in merged_data]
    result = pd.DataFrame(rows, columns=unified_cols)

    # Đánh lại STT từ 1
    result["STT"] = range(1, len(result) + 1)

    return result.reset_index(drop=True)


# ══════════════════════════════════════════════════════════════════════════════
# XUẤT FILE VÀ ĐỊNH DẠNG
# ══════════════════════════════════════════════════════════════════════════════

def apply_format(excel_path: str, max_trich: int):
    """Định dạng file Excel đầu ra: header, màu sắc, căn lề, kích thước cột."""
    print("🎨 Đang áp dụng định dạng Excel...")
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active

    thin = Border(
        left=Side(border_style="thin", color="CCCCCC"),
        right=Side(border_style="thin", color="CCCCCC"),
        top=Side(border_style="thin", color="CCCCCC"),
        bottom=Side(border_style="thin", color="CCCCCC"),
    )

    # Header row
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor="1F3864")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin
    ws.row_dimensions[1].height = 40

    # Data rows: màu xen kẽ
    for row_idx in range(2, ws.max_row + 1):
        bg = "FFFFFF" if row_idx % 2 == 0 else "F0F4FA"
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.fill = PatternFill("solid", fgColor=bg)
            cell.border = thin
            is_center_col = col_idx in (1, 6)
            cell.alignment = Alignment(
                horizontal="center" if is_center_col else "left",
                vertical="top",
                wrap_text=True,
            )
            # Cột "Trích nguồn" → in nghiêng, cỡ nhỏ hơn
            is_trich_nguon = col_idx > 6 and col_idx % 2 == 0
            cell.font = Font(italic=is_trich_nguon, size=9 if col_idx > 6 else 10)

    # Độ rộng cột
    col_widths = {1: 5, 2: 30, 3: 36, 4: 38, 5: 38, 6: 18}
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        if col_idx in col_widths:
            ws.column_dimensions[letter].width = col_widths[col_idx]
        else:
            # Cột nội dung trích dẫn (lẻ) rộng hơn cột trích nguồn (chẵn)
            ws.column_dimensions[letter].width = 65 if col_idx % 2 != 0 else 45

    ws.freeze_panes = "A2"
    wb.save(excel_path)
    print("✨ Định dạng hoàn tất!")


# ══════════════════════════════════════════════════════════════════════════════
# MAIN
# ══════════════════════════════════════════════════════════════════════════════

def parse_args():
    parser = argparse.ArgumentParser(
        description="Gộp nhiều file Excel Lịch Sử 11 thành một file thống nhất",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Ví dụ sử dụng:
  # Lần đầu gộp 3 file:
  python merge_lichsu.py --files A.xlsx B.xlsx C.xlsx

  # Thêm file mới vào file lớn đã có (không ghi đè):
  python merge_lichsu.py --base LichSu11.xlsx --files D.xlsx

  # Gộp tất cả .xlsx trong thư mục pending/:
  python merge_lichsu.py --folder ./pending/ --base LichSu11.xlsx

  # Tuỳ chỉnh tên file đầu ra:
  python merge_lichsu.py --files A.xlsx B.xlsx --output TongHop.xlsx

  # Kết hợp folder + base + output:
  python merge_lichsu.py --base LichSu11.xlsx --folder ./them_moi/ --output LichSu11_v2.xlsx
        """,
    )
    parser.add_argument(
        "--files", nargs="+", metavar="FILE",
        help="Danh sách file .xlsx cần gộp (tách bằng dấu cách)",
    )
    parser.add_argument(
        "--folder", metavar="FOLDER",
        help="Thư mục chứa file .xlsx cần gộp (sẽ đọc tất cả .xlsx trong thư mục)",
    )
    parser.add_argument(
        "--base", metavar="BASE_FILE",
        help="File lớn đã có (nếu có). Nội dung file này KHÔNG bị ghi đè. "
             "Các file --files/--folder sẽ được thêm vào sau.",
    )
    parser.add_argument(
        "--output", metavar="OUTPUT", default="LichSu11.xlsx",
        help="Tên file đầu ra (mặc định: LichSu11.xlsx)",
    )
    parser.add_argument(
        "--no-format", action="store_true",
        help="Bỏ qua bước định dạng Excel (nhanh hơn, dùng khi cần debug)",
    )
    return parser.parse_args()


def collect_input_files(args) -> tuple[str | None, list[str]]:
    """
    Trả về (base_file, [danh_sach_file_can_gop]).
    base_file: file lớn đã có (không gộp lại, chỉ dùng làm nền)
    danh_sach_file_can_gop: các file sẽ gộp thêm vào
    """
    base_file: str | None = None
    new_files: list[str] = []

    # File base (đã tồn tại, không ghi đè)
    if args.base:
        if not os.path.exists(args.base):
            print(f"❌ Không tìm thấy file base: {args.base}")
            sys.exit(1)
        base_file = args.base
        print(f"📌 File base (nền): {base_file}")

    # Files chỉ định trực tiếp
    if args.files:
        for f in args.files:
            if not os.path.exists(f):
                print(f"⚠️  Không tìm thấy file: {f} → bỏ qua")
            else:
                # Không thêm file base vào danh sách gộp để tránh trùng
                if base_file and os.path.abspath(f) == os.path.abspath(base_file):
                    print(f"  ℹ️  Bỏ qua (đã là base): {f}")
                else:
                    new_files.append(f)

    # Files từ folder
    if args.folder:
        if not os.path.isdir(args.folder):
            print(f"❌ Không tìm thấy thư mục: {args.folder}")
            sys.exit(1)
        folder_files = sorted(glob.glob(os.path.join(args.folder, "*.xlsx")))
        for f in folder_files:
            if base_file and os.path.abspath(f) == os.path.abspath(base_file):
                print(f"  ℹ️  Bỏ qua (đã là base): {f}")
                continue
            if f not in new_files:
                new_files.append(f)
        print(f"📂 Tìm thấy {len(folder_files)} file trong folder: {args.folder}")

    if not base_file and not new_files:
        print("❌ Chưa chỉ định file nào để gộp!")
        print("   Dùng --files hoặc --folder để chỉ định file đầu vào.")
        sys.exit(1)

    if not new_files and base_file:
        print("⚠️  Không có file mới nào để gộp vào base. Sẽ sao chép base → output.")

    return base_file, new_files


def main():
    args = parse_args()

    print("=" * 60)
    print("  GỘP FILE EXCEL LỊCH SỬ 11")
    print(f"  Thời gian: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    print("=" * 60)

    base_file, new_files = collect_input_files(args)

    # ── Đọc tất cả file ──────────────────────────────────────────────────────
    print("\n📖 Đang đọc file...")
    base_df: pd.DataFrame | None = None
    if base_file:
        print(f"  [BASE] {base_file}")
        base_df = read_excel_safe(base_file)

    new_dfs: list[pd.DataFrame] = []
    for path in new_files:
        print(f"  [NEW]  {path}")
        new_dfs.append(read_excel_safe(path))

    # ── Xây dựng cột thống nhất ───────────────────────────────────────────────
    all_dfs = ([base_df] if base_df is not None else []) + new_dfs
    if not all_dfs:
        print("❌ Không có dữ liệu nào để xử lý!")
        sys.exit(1)

    unified_cols = build_unified_columns(all_dfs)
    max_trich = (len(unified_cols) - len(FIXED_COLS)) // 2
    print(f"\n📐 Cấu trúc cột thống nhất: {len(FIXED_COLS)} cột cố định + {max_trich} cặp trích dẫn = {len(unified_cols)} cột tổng")

    # ── Căn chỉnh tất cả DataFrame về cùng cột ────────────────────────────────
    print("\n🔧 Căn chỉnh cột...")
    if base_df is not None:
        base_df = align_df_to_columns(base_df, unified_cols)
    new_dfs = [align_df_to_columns(df, unified_cols) for df in new_dfs]

    # ── Gộp dữ liệu ──────────────────────────────────────────────────────────
    print("\n🔗 Đang gộp dữ liệu...")
    if base_df is not None:
        print(f"  Base có: {len(base_df)} dòng")
    for i, (df, path) in enumerate(zip(new_dfs, new_files)):
        print(f"  File {i+1} ({Path(path).name}): {len(df)} dòng")

    result_df = merge_dataframes(base_df, new_dfs, unified_cols)
    print(f"\n  ✅ Kết quả: {len(result_df)} dòng tổng cộng")

    # ── Lưu file ─────────────────────────────────────────────────────────────
    output_path = args.output
    print(f"\n💾 Đang lưu → {output_path}...")
    result_df.to_excel(output_path, index=False)
    print(f"  ✅ Đã lưu file ({os.path.getsize(output_path) / 1024:.1f} KB)")

    # ── Định dạng ────────────────────────────────────────────────────────────
    if not args.no_format:
        apply_format(output_path, max_trich)

    # ── Tổng kết ─────────────────────────────────────────────────────────────
    print(f"""
╔══════════════════════════════════════════════════════╗
║  ✅ HOÀN TẤT GỘP FILE
║  📄 Output : {output_path}
║  📊 Tổng số dòng dữ liệu : {len(result_df)}
║  📋 Tổng số cột          : {len(unified_cols)}
║  🗂️  Cặp trích dẫn tối đa : {max_trich}
╚══════════════════════════════════════════════════════╝
""")


if __name__ == "__main__":
    main()