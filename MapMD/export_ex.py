"""
export_beautiful.py — Xuất Excel màu sắc đẹp + Có cột đếm "Số lượng đoạn trích"
"""

import json
import os
import sys
import argparse
from collections import defaultdict

try:
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.worksheet import Worksheet
except ImportError:
    print("❌ Thiếu openpyxl. Chạy lệnh: pip install openpyxl")
    sys.exit(1)

# ============================================================
# CONSTANTS — MÀU SẮC & STYLE 
# ============================================================
COLOR_HEADER_BG      = "1F3864"   # Navy đậm
COLOR_HEADER_FG      = "FFFFFF"   # Trắng
COLOR_ZEBRA_ODD      = "FFFFFF"   # Hàng lẻ: trắng
COLOR_ZEBRA_EVEN     = "F5F7FA"   # Hàng chẵn: xám rất nhạt
COLOR_TIEU_MUC_BG    = "EBF3FB"   # Xanh nhạt — cột Tiểu mục có nội dung
COLOR_TIEU_MUC_EMPTY = "FAFAFA"   # Xám rất nhạt — Tiểu mục trống
COLOR_NOI_DUNG_BG    = "FFFDE7"   # Vàng nhạt — cột Nội dung
COLOR_TRICH_NGUON_BG = "E8F5E9"   # Xanh lá nhạt — cột Trích nguồn
COLOR_COUNT_BG       = "FFF3E0"   # Cam nhạt - cột Đếm số lượng đoạn trích

COL_STT        = 1
COL_CHUONG     = 2
COL_BAI        = 3
COL_MUC        = 4
COL_TIEU_MUC   = 5
COL_SO_DOAN    = 6   # CỘT MỚI: Số lượng đoạn trích
N_FIXED        = 6   # Tăng số cột cố định lên 6

def _thin_border() -> Border:
    thin = Side(border_style="thin", color="CCCCCC")
    return Border(left=thin, right=thin, top=thin, bottom=thin)

def load_json(path):
    with open(path, 'r', encoding='utf-8') as f:
        return json.load(f)

# ============================================================
# CLASS ExcelBeautifulExporter
# ============================================================
class ExcelBeautifulExporter:
    def __init__(self):
        self.wb = openpyxl.Workbook()
        self._rows = []
        self._max_pairs = 0

    def process_data(self, toc_path, map_path):
        """Đọc và ghép nối file toc.json và fast_md_content_map.json"""
        toc_data = load_json(toc_path)
        map_data = load_json(map_path)

        # 1. Gom dữ liệu Map theo cache_key
        grouped_excerpts = defaultdict(list)
        for item in map_data:
            for mid in item.get("mapping_ids", []):
                grouped_excerpts[mid].append({
                    "noi_dung": item.get("excerpt", ""),
                    "trich_nguon": item.get("trich_nguon", "")
                })

        # 2. Xây dựng Rows từ TOC
        chapters = toc_data.get("chapters", toc_data) if isinstance(toc_data, dict) else toc_data
        for ch in chapters:
            ch_num = ch.get("chapter_number", "")
            ch_title = ch.get("chapter_title", "")
            for ls in ch.get("lessons", []):
                ls_num = ls.get("lesson_number", "")
                ls_title = ls.get("lesson_title", "")
                bai_str = f"Bài {ls_num}. {ls_title}" if ls_num else ls_title
                for sc in ls.get("sections", []):
                    sc_num = sc.get("section_number", "")
                    sc_title = sc.get("section_title", "")
                    subsecs = sc.get("subsections", [])
                    
                    if subsecs:
                        for sub in subsecs:
                            sub_title = sub.get("subsection_title", "")
                            cache_key = f"{ch_num}|{ls_num}|{sc_num}|{sub_title}"
                            trich_doan = grouped_excerpts.get(cache_key, [])
                            self._add_row(ch_title, bai_str, sc_title, sub_title, trich_doan)
                    else:
                        cache_key = f"{ch_num}|{ls_num}|{sc_num}|"
                        trich_doan = grouped_excerpts.get(cache_key, [])
                        self._add_row(ch_title, bai_str, sc_title, "", trich_doan)

    def _add_row(self, ch_title, bai_str, sc_title, sub_title, trich_doan):
        if len(trich_doan) > self._max_pairs:
            self._max_pairs = len(trich_doan)
        self._rows.append({
            "chuong": ch_title,
            "bai": bai_str,
            "muc": sc_title,
            "tieu_muc": sub_title,
            "trich_doan": trich_doan,
        })

    def export(self, output_path):
        """Tạo các sheet và xuất file"""
        ws_main = self.wb.active
        ws_main.title = "Trích dẫn"
        self._build_main_sheet(ws_main)

        ws_summary = self.wb.create_sheet("Thống kê")
        self._build_summary_sheet(ws_summary)

        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        self.wb.save(output_path)
        print(f"💾 Đã lưu Excel thành công tại: {output_path}")

    def _build_main_sheet(self, ws: Worksheet):
        N_PAIRS = max(1, self._max_pairs)
        TOTAL_COLS = N_FIXED + N_PAIRS * 2

        # Header (Thêm cột Số lượng đoạn trích)
        all_headers = ["STT", "Chủ đề / Chương", "Bài", "Mục", "Tiểu mục", "Số lượng đoạn trích"]
        for i in range(1, N_PAIRS + 1):
            all_headers.extend([f"Nội dung trích dẫn {i}", f"Trích nguồn {i}"])

        for col_idx, header in enumerate(all_headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = Font(bold=True, color=COLOR_HEADER_FG, size=11)
            cell.fill = PatternFill("solid", fgColor=COLOR_HEADER_BG)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = _thin_border()
        ws.row_dimensions[1].height = 36

        # Data Rows
        for stt, row_data in enumerate(self._rows, start=1):
            data_row = stt + 1
            trich_doan = row_data["trich_doan"]
            has_sub = bool(row_data["tieu_muc"])
            zebra = COLOR_ZEBRA_ODD if stt % 2 == 1 else COLOR_ZEBRA_EVEN
            
            so_luong_doan = len(trich_doan)

            # Cập nhật mảng cố định để đưa số lượng đoạn trích vào
            fixed_values = [
                stt, 
                row_data["chuong"], 
                row_data["bai"], 
                row_data["muc"], 
                row_data["tieu_muc"],
                so_luong_doan # CỘT ĐẾM
            ]
            
            for col_idx, val in enumerate(fixed_values, start=1):
                cell = ws.cell(row=data_row, column=col_idx, value=val or "")
                
                # Tô màu riêng biệt
                if col_idx == COL_TIEU_MUC:
                    bg = COLOR_TIEU_MUC_BG if has_sub else COLOR_TIEU_MUC_EMPTY
                elif col_idx == COL_SO_DOAN:
                    bg = COLOR_COUNT_BG if so_luong_doan > 0 else COLOR_ZEBRA_EVEN
                else:
                    bg = zebra
                    
                cell.fill = PatternFill("solid", fgColor=bg)
                cell.border = _thin_border()
                
                # Căn giữa cho STT và Số lượng đoạn trích
                is_center = col_idx in (COL_STT, COL_SO_DOAN)
                cell.alignment = Alignment(horizontal="center" if is_center else "left", vertical="top", wrap_text=True)
                
                # Làm đậm tên Bài và Số lượng đoạn trích (nếu > 0)
                is_bold = (col_idx == COL_BAI) or (col_idx == COL_SO_DOAN and so_luong_doan > 0)
                cell.font = Font(bold=is_bold, size=10, color="FF0000" if col_idx == COL_SO_DOAN and so_luong_doan > 0 else "000000")

            # Columns Nội dung / Trích nguồn
            for pair_idx in range(N_PAIRS):
                col_nd = N_FIXED + pair_idx * 2 + 1
                col_ts = N_FIXED + pair_idx * 2 + 2

                if pair_idx < len(trich_doan):
                    nd = trich_doan[pair_idx].get("noi_dung", "")
                    ts = trich_doan[pair_idx].get("trich_nguon", "")
                else:
                    nd, ts = "", ""

                cell_nd = ws.cell(row=data_row, column=col_nd, value=nd)
                cell_nd.fill = PatternFill("solid", fgColor=COLOR_NOI_DUNG_BG if nd else zebra)
                cell_nd.border = _thin_border()
                cell_nd.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                cell_nd.font = Font(size=10)

                cell_ts = ws.cell(row=data_row, column=col_ts, value=ts)
                cell_ts.fill = PatternFill("solid", fgColor=COLOR_TRICH_NGUON_BG if ts else zebra)
                cell_ts.border = _thin_border()
                cell_ts.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                cell_ts.font = Font(italic=True, size=9)

            # Auto Row Height
            if trich_doan:
                max_nd_len = max((len(str(item.get("noi_dung", ""))) for item in trich_doan), default=0)
                row_height = max(25, min(300, max_nd_len // 12 * 15))
            else:
                row_height = 25
            ws.row_dimensions[data_row].height = row_height

        # Column Widths
        col_widths_fixed = {
            COL_STT: 5, 
            COL_CHUONG: 28, 
            COL_BAI: 35, 
            COL_MUC: 38, 
            COL_TIEU_MUC: 38,
            COL_SO_DOAN: 18 # Chiều rộng cột đếm
        }
        for col_idx, width in col_widths_fixed.items():
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        for pair_idx in range(N_PAIRS):
            ws.column_dimensions[get_column_letter(N_FIXED + pair_idx * 2 + 1)].width = 60
            ws.column_dimensions[get_column_letter(N_FIXED + pair_idx * 2 + 2)].width = 42

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(TOTAL_COLS)}{len(self._rows) + 1}"

    def _build_summary_sheet(self, ws: Worksheet):
        ws.column_dimensions["A"].width = 42
        ws.column_dimensions["B"].width = 15
        hdr_fill = PatternFill("solid", fgColor=COLOR_HEADER_BG)
        hdr_font = Font(bold=True, color=COLOR_HEADER_FG, size=11)

        def dat(r, c, val, bold=False):
            cell = ws.cell(row=r, column=c, value=val)
            cell.font = Font(bold=bold, size=10)
            cell.border = _thin_border()
            cell.alignment = Alignment(vertical="center", wrap_text=True)

        ws.cell(1, 1, "Thống kê").fill = hdr_fill
        ws.cell(1, 1).font = hdr_font
        ws.cell(1, 2, "Số lượng").fill = hdr_fill
        ws.cell(1, 2).font = hdr_font
        ws.row_dimensions[1].height = 28

        row = 2
        total_nodes = len(self._rows)
        has_sub = sum(1 for r in self._rows if r.get("tieu_muc"))
        total_td = sum(len(r["trich_doan"]) for r in self._rows)
        total_rong = sum(1 for r in self._rows if not r["trich_doan"])

        stats = [
            ("Tổng số node lá (mục/tiểu mục)", total_nodes, True),
            ("  - Có tiểu mục (subsection)", has_sub, False),
            ("  - Chỉ có mục (section)", total_nodes - has_sub, False),
            ("Tổng đoạn trích dẫn", total_td, True),
            ("  - Node không có đoạn trích (rỗng)", total_rong, False),
            ("Số cặp nội dung/trích nguồn tối đa / node", self._max_pairs, False),
        ]

        for label, val, is_bold in stats:
            dat(row, 1, label, bold=is_bold)
            dat(row, 2, val)
            row += 1

        row += 1
        ws.cell(row, 1, "Chủ đề / Chương").fill = hdr_fill
        ws.cell(row, 1).font = hdr_font
        ws.cell(row, 2, "Số node").fill = hdr_fill
        ws.cell(row, 2).font = hdr_font
        row += 1

        from collections import Counter
        chuong_counter = Counter(r["chuong"] for r in self._rows)
        for chuong, count in sorted(chuong_counter.items()):
            dat(row, 1, chuong)
            dat(row, 2, count)
            row += 1

def main():
    parser = argparse.ArgumentParser(description="Xuất kết quả Map ra Excel (Màu sắc theo mẫu + Đếm số lượng)")
    parser.add_argument("--toc", required=True, help="Đường dẫn file toc.json")
    parser.add_argument("--map", required=True, help="Đường dẫn file fast_md_content_map.json")
    parser.add_argument("--output", default="KetQua_Map_Dep.xlsx", help="Tên file Excel đầu ra")
    args = parser.parse_args()

    print("=" * 60)
    print("🎨 ĐANG TẠO FILE EXCEL (CÓ ĐẾM SỐ LƯỢNG ĐOẠN TRÍCH)...")
    
    exporter = ExcelBeautifulExporter()
    try:
        exporter.process_data(args.toc, args.map)
        exporter.export(args.output)
    except Exception as e:
        print(f"❌ Lỗi: {e}")

if __name__ == "__main__":
    main()